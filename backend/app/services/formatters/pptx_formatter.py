"""
PPTX Formatter Service (Stage 6).
Renders presentation slide deck artefacts into 16:9 widescreen PowerPoint (.pptx) files with speaker notes.
"""

from pathlib import Path
from typing import Dict, Any, List, Optional
from backend.app.core.logging import logger
from backend.app.schemas.generated_artefact import GeneratedArtefact


class PptxFormatter:
    """
    Constructs styled .pptx slide decks from PresentationDeckContent payload using python-pptx.
    """

    def generate_pptx(self, artefact: GeneratedArtefact, output_path: Path) -> Path:
        """
        Builds a .pptx file and saves it to output_path.
        """
        output_path.parent.mkdir(parents=True, exist_ok=True)
        content = artefact.content or {}

        try:
            import pptx
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor

            prs = pptx.Presentation()
            # Set 16:9 Widescreen dimensions
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)

            blank_layout = prs.slide_layouts[6]  # Blank slide

            # 1. Slide 1: Title Slide
            title_slide = prs.slides.add_slide(blank_layout)
            
            # Title Box
            tb_title = title_slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(2.0))
            tf_title = tb_title.text_frame
            tf_title.word_wrap = True
            p_title = tf_title.paragraphs[0]
            p_title.text = content.get("presentation_title", content.get("title", "Executive Presentation Deck"))
            p_title.font.bold = True
            p_title.font.size = Pt(36)
            p_title.font.color.rgb = RGBColor(20, 35, 60)

            # Subtitle Box
            tagline = content.get("tagline", content.get("hook", content.get("overview", "")))
            if tagline:
                p_sub = tf_title.add_paragraph()
                p_sub.text = tagline[:140]
                p_sub.font.size = Pt(20)
                p_sub.font.italic = True
                p_sub.font.color.rgb = RGBColor(100, 110, 125)

            # 2. Content Slides
            slides_data = content.get("slides", [])
            if not slides_data and "key_findings" in content:
                # Synthesize slide list if generic content
                slides_data = [
                    {"title": "Key Findings", "bullet_points": content.get("key_findings", []), "speaker_notes": "Review findings."},
                    {"title": "Recommendations", "bullet_points": content.get("recommendations", []), "speaker_notes": "Review action plan."}
                ]

            for idx, slide_info in enumerate(slides_data, start=2):
                if not isinstance(slide_info, dict):
                    continue

                slide = prs.slides.add_slide(blank_layout)

                # Slide Header Title
                tb_head = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0))
                tf_head = tb_head.text_frame
                tf_head.word_wrap = True
                p_head = tf_head.paragraphs[0]
                p_head.text = slide_info.get("title", f"Slide {idx}")
                p_head.font.bold = True
                p_head.font.size = Pt(28)
                p_head.font.color.rgb = RGBColor(20, 35, 60)

                # Left Column: Bullet Points
                tb_bullets = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(7.5), Inches(4.8))
                tf_bullets = tb_bullets.text_frame
                tf_bullets.word_wrap = True

                bullets = slide_info.get("bullet_points", [])
                for b_idx, bullet in enumerate(bullets):
                    p_b = tf_bullets.paragraphs[0] if b_idx == 0 else tf_bullets.add_paragraph()
                    p_b.text = f"•  {bullet}"
                    p_b.font.size = Pt(18)
                    p_b.font.color.rgb = RGBColor(40, 40, 40)
                    p_b.space_after = Pt(14)

                # Right Column: Visual Recommendation Box
                vis = slide_info.get("visual_suggestion") or slide_info.get("visual_cue") or "Suggested Layout: Data Callout Card"
                tb_vis = slide.shapes.add_textbox(Inches(8.6), Inches(1.8), Inches(3.9), Inches(4.8))
                tf_vis = tb_vis.text_frame
                tf_vis.word_wrap = True
                p_vis_head = tf_vis.paragraphs[0]
                p_vis_head.text = "[ VISUAL ASSET BLUEPRINT ]"
                p_vis_head.font.bold = True
                p_vis_head.font.size = Pt(12)
                p_vis_head.font.color.rgb = RGBColor(0, 102, 204)

                p_vis_body = tf_vis.add_paragraph()
                p_vis_body.text = vis
                p_vis_body.font.size = Pt(14)
                p_vis_body.font.italic = True
                p_vis_body.font.color.rgb = RGBColor(80, 80, 80)

                # Attach Speaker Notes
                notes = slide_info.get("speaker_notes", "")
                if notes and hasattr(slide, "notes_slide"):
                    notes_slide = slide.notes_slide
                    notes_slide.notes_text_frame.text = f"Speaker Notes:\n{notes}"

            prs.save(str(output_path))
            logger.info("Generated PPTX deliverable at %s", output_path)
            return output_path

        except Exception as exc:
            logger.warning("python-pptx rendering failed (%s); writing raw fallback text file", exc)
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(f"PRESENTATION SLIDE DECK: {content.get('presentation_title', 'Deck')}\n\n")
                for s in content.get("slides", []):
                    f.write(f"Slide: {s.get('title')}\n")
                    for b in s.get("bullet_points", []):
                        f.write(f"  - {b}\n")
                    f.write(f"Notes: {s.get('speaker_notes')}\n\n")
            return output_path


pptx_formatter = PptxFormatter()
